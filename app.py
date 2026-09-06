import os
import json
import re

import streamlit as st
from pypdf import PdfReader
from docx import Document
from google import genai
from google.genai import types


# -----------------------------
# Page setup
# -----------------------------
st.set_page_config(
    page_title="AI Resume ATS Checker",
    page_icon="📄",
    layout="wide",
)

st.title("📄 AI Resume ATS Checker")
st.write(
    "Upload your resume and optionally add a job description. "
    "The app will estimate an ATS score and suggest practical improvements."
)

# -----------------------------
# Helpers
# -----------------------------
def get_api_key():
    """Get the Gemini API key from Streamlit Secrets or environment variables."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
        if key:
            return key
    except Exception:
        pass

    return os.getenv("GEMINI_API_KEY")


def extract_pdf_text(uploaded_file):
    """Extract text from a PDF."""
    reader = PdfReader(uploaded_file)
    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)

    return "\n".join(pages).strip()


def extract_docx_text(uploaded_file):
    """Extract text from a DOCX file."""
    document = Document(uploaded_file)

    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Also read text from tables because some resumes use tables.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    return "\n".join(paragraphs).strip()


def extract_resume_text(uploaded_file):
    """Extract text according to the uploaded file type."""
    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):
        return extract_pdf_text(uploaded_file)

    if file_name.endswith(".docx"):
        return extract_docx_text(uploaded_file)

    raise ValueError("Only PDF and DOCX files are supported.")


def clean_score(value):
    """Keep a score between 0 and 100."""
    try:
        return max(0, min(100, int(value)))
    except (TypeError, ValueError):
        return 0


# -----------------------------
# Gemini configuration
# -----------------------------
# Keep the model in one place so it can easily be changed later.
MODEL_NAME = "gemini-3.6-flash"

RESULT_SCHEMA = {
    "type": "object",
    "properties": {
        "ats_score": {
            "type": "integer",
            "description": "Overall ATS score from 0 to 100."
        },
        "keyword_score": {
            "type": "integer",
            "description": "Keyword/job-description match score from 0 to 100."
        },
        "formatting_score": {
            "type": "integer",
            "description": "ATS-friendly formatting score from 0 to 100."
        },
        "section_score": {
            "type": "integer",
            "description": "Resume section completeness score from 0 to 100."
        },
        "experience_score": {
            "type": "integer",
            "description": "Relevance and strength of experience score from 0 to 100."
        },
        "summary": {
            "type": "string",
            "description": "Short explanation of the overall result."
        },
        "matched_keywords": {
            "type": "array",
            "items": {"type": "string"},
            "description": "Important keywords found in both resume and job description."
        },
        "missing_keywords": {
            "type": "array",
            "items": {"type": "string"},
            "description": "Important job-description keywords missing from the resume."
        },
        "strengths": {
            "type": "array",
            "items": {"type": "string"},
            "description": "Specific strengths of the resume."
        },
        "improvements": {
            "type": "array",
            "items": {"type": "string"},
            "description": "Specific and actionable improvements."
        },
        "formatting_warnings": {
            "type": "array",
            "items": {"type": "string"},
            "description": "Possible ATS formatting problems."
        }
    },
    "required": [
        "ats_score",
        "keyword_score",
        "formatting_score",
        "section_score",
        "experience_score",
        "summary",
        "matched_keywords",
        "missing_keywords",
        "strengths",
        "improvements",
        "formatting_warnings"
    ]
}


def analyze_resume(resume_text, job_description):
    """Send the resume to Gemini and return structured ATS feedback."""
    api_key = get_api_key()

    if not api_key:
        raise ValueError(
            "Gemini API key was not found. Add GEMINI_API_KEY to "
            "Streamlit Secrets or your environment variables."
        )

    client = genai.Client(api_key=api_key)

    if job_description.strip():
        job_text = job_description.strip()
    else:
        job_text = (
            "No job description was provided. Evaluate the resume using "
            "general ATS-friendly resume standards and identify broadly "
            "useful skills/keywords."
        )

    prompt = f"""
You are an ATS resume evaluator.

Analyze the resume below.

IMPORTANT:
1. Give an ATS-style score from 0 to 100.
2. Do not claim that your score is the exact score of a real ATS. It is an estimate.
3. If a job description is provided, compare the resume against it.
4. If no job description is provided, evaluate general ATS quality.
5. Do not invent experience, education, certifications, skills, or keywords.
6. Only recommend adding a keyword if it is genuinely relevant to the target job.
7. Consider common ATS issues such as:
   - unusual formatting
   - tables/text boxes
   - missing standard sections
   - unclear job titles
   - weak action verbs
   - lack of measurable achievements
   - missing relevant keywords
   - excessive graphics/symbols
   - poor readability
8. Make the improvements specific and actionable.
9. Return only the requested structured JSON.

JOB DESCRIPTION:
{job_text}

RESUME:
{resume_text}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=RESULT_SCHEMA,
            temperature=0.2,
        ),
    )

    if not response.text:
        raise ValueError("Gemini returned an empty response.")

    return json.loads(response.text)


# -----------------------------
# User interface
# -----------------------------
uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"],
    help="PDF or DOCX only.",
)

job_description = st.text_area(
    "Job Description (optional)",
    height=220,
    placeholder="Paste the job description here for a more accurate ATS match.",
)

analyze_button = st.button(
    "🔍 Analyze Resume",
    type="primary",
    use_container_width=True,
)

if analyze_button:
    if uploaded_file is None:
        st.warning("Please upload a PDF or DOCX resume first.")
        st.stop()

    try:
        with st.spinner("Reading your resume..."):
            resume_text = extract_resume_text(uploaded_file)

        if len(resume_text.strip()) < 100:
            st.error(
                "Very little text could be extracted from this file. "
                "If this is a scanned/image-only PDF, please upload a text-based "
                "PDF or DOCX version."
            )
            st.stop()

        # Prevent unnecessarily huge prompts.
        resume_text = resume_text[:50000]
        job_description = job_description[:30000]

        with st.spinner("Gemini is analyzing your resume..."):
            result = analyze_resume(resume_text, job_description)

        # -----------------------------
        # Score cards
        # -----------------------------
        st.subheader("📊 ATS Results")

        score_cols = st.columns(5)

        scores = [
            ("ATS Score", result["ats_score"]),
            ("Keywords", result["keyword_score"]),
            ("Formatting", result["formatting_score"]),
            ("Sections", result["section_score"]),
            ("Experience", result["experience_score"]),
        ]

        for col, (label, value) in zip(score_cols, scores):
            with col:
                st.metric(label, f"{clean_score(value)}/100")

        st.progress(clean_score(result["ats_score"]) / 100)

        st.subheader("📝 Overall Assessment")
        st.write(result["summary"])

        # -----------------------------
        # Keywords
        # -----------------------------
        col1, col2 = st.columns(2)

        with col1:
            st.subheader("✅ Matched Keywords")
            matched = result.get("matched_keywords", [])
            if matched:
                for item in matched:
                    st.write(f"• {item}")
            else:
                st.info("No important matched keywords were identified.")

        with col2:
            st.subheader("⚠️ Missing Keywords")
            missing = result.get("missing_keywords", [])
            if missing:
                for item in missing:
                    st.write(f"• {item}")
            else:
                st.success("No major missing keywords were identified.")

        # -----------------------------
        # Strengths & improvements
        # -----------------------------
        col1, col2 = st.columns(2)

        with col1:
            st.subheader("💪 Strengths")
            for item in result.get("strengths", []):
                st.write(f"• {item}")

        with col2:
            st.subheader("🚀 Improvements")
            for item in result.get("improvements", []):
                st.write(f"• {item}")

        # -----------------------------
        # Formatting warnings
        # -----------------------------
        st.subheader("📐 ATS Formatting Warnings")

        warnings = result.get("formatting_warnings", [])

        if warnings:
            for warning in warnings:
                st.warning(warning)
        else:
            st.success("No major ATS formatting problems were identified.")

        # -----------------------------
        # Simple downloadable report
        # -----------------------------
        report = {
            "ATS Score": clean_score(result["ats_score"]),
            "Keyword Score": clean_score(result["keyword_score"]),
            "Formatting Score": clean_score(result["formatting_score"]),
            "Section Score": clean_score(result["section_score"]),
            "Experience Score": clean_score(result["experience_score"]),
            "Summary": result["summary"],
            "Matched Keywords": result.get("matched_keywords", []),
            "Missing Keywords": result.get("missing_keywords", []),
            "Strengths": result.get("strengths", []),
            "Improvements": result.get("improvements", []),
            "Formatting Warnings": result.get("formatting_warnings", []),
        }

        st.download_button(
            "⬇️ Download Analysis",
            data=json.dumps(report, indent=2),
            file_name="resume_ats_analysis.json",
            mime="application/json",
            use_container_width=True,
        )

    except Exception as error:
        st.error("The resume could not be analyzed.")
        st.exception(error)
else:
    st.info(
        "Upload a resume and click **Analyze Resume**. "
        "Adding a job description gives a more useful keyword-match score."
    )
